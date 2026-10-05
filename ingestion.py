import io
import pymupdf as fitz
from docx import Document
from typing import List, Dict, Any

def parse_document(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    if not file_bytes or len(file_bytes) == 0:
        raise ValueError(f"File '{filename}' is empty (0 bytes).")

    ext = filename.lower().split(".")[-1]
    pages = []

    if ext == "pdf":
        try:
            with fitz.open(stream=file_bytes, filetype="pdf") as doc:
                if doc.is_encrypted:
                    raise PermissionError(f"PDF '{filename}' is password protected.")
                if len(doc) == 0:
                    raise ValueError(f"PDF '{filename}' contains 0 pages.")

                for page_idx, page in enumerate(doc):
                    text = page.get_text("text").strip()
                    if text:
                        pages.append({
                            "page_number": page_idx + 1,
                            "page": page_idx + 1,
                            "text": text,
                            "filename": filename
                        })
        except (PermissionError, ValueError):
            raise
        except Exception as e:
            raise RuntimeError(f"Failed to read PDF '{filename}': {str(e)}")

    elif ext == "docx":
        try:
            doc = Document(io.BytesIO(file_bytes))
            full_text = []
            for para in doc.paragraphs:
                cleaned = para.text.strip()
                if cleaned:
                    full_text.append(cleaned)

            # Also extract tables commonly found in syllabus and regulations
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        # Deduplicate repeated adjacent merged cells
                        deduped_cells = []
                        for cell_val in row_cells:
                            if not deduped_cells or deduped_cells[-1] != cell_val:
                                deduped_cells.append(cell_val)
                        if deduped_cells:
                            full_text.append(" | ".join(deduped_cells))

            combined_text = "\n".join(full_text)
            if combined_text:
                pages.append({
                    "page_number": 1,
                    "page": 1,
                    "text": combined_text,
                    "filename": filename
                })
        except Exception as e:
            raise RuntimeError(f"Failed to parse DOCX '{filename}': {str(e)}")

    elif ext in ["txt", "md", "csv"]:
        try:
            # Handle UTF-8 with fallback to Latin-1
            try:
                raw_text = file_bytes.decode("utf-8").strip()
            except UnicodeDecodeError:
                raw_text = file_bytes.decode("latin-1", errors="replace").strip()

            if raw_text:
                pages.append({
                    "page_number": 1,
                    "page": 1,
                    "text": raw_text,
                    "filename": filename
                })
        except Exception as e:
            raise RuntimeError(f"Failed to decode text file '{filename}': {str(e)}")
    else:
        raise ValueError(f"Unsupported file format: '.{ext}'")

    if not pages:
        raise ValueError(f"No extractable text found in '{filename}'. Scanned images require OCR.")

    return pages

def extract_from_txt(file_bytes: bytes, filename: str = "document.txt") -> List[Dict[str, Any]]:
    """Helper extraction function for plain text files."""
    return parse_document(file_bytes, filename)